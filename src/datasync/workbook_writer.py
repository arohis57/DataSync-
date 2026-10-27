"""Write a DataSync workbook (Excel tables with id-based links) from a cockpit payload.

The reverse of excel_to_model. Used to build test/sample workbooks, and a reference for the
ADO-to-Excel script, which must produce the same sheets, tables and headers.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .schema import (
    ACTION_LINKS,
    BURNUP_HEADERS,
    BURNUP_SHEET,
    BURNUP_TABLE,
    META_KEYS,
    META_SHEET,
    META_TABLE,
    SHEETS,
)


def add_table(wb: Workbook, sheet: str, table: str, headers: list[str], rows: list[list], text_cols: set[int] = frozenset(), date_cols: set[int] = frozenset()) -> None:
    ws = wb.create_sheet(sheet)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = max(12, min(40, len(h) + 4))
    for r in range(2, len(rows) + 2):
        for c in text_cols:  # keep ids as text so Excel doesn't turn them into numbers
            ws.cell(r, c + 1).number_format = "@"
        for c in date_cols:
            if ws.cell(r, c + 1).value is not None:
                ws.cell(r, c + 1).number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    tbl = Table(displayName=table, ref=f"A1:{get_column_letter(len(headers))}{len(rows) + 1}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)


def write_workbook(payload: dict, out: Path) -> None:
    data = payload["data"]
    key_pos = {s.key: next(i for i, c in enumerate(s.cols) if c.kind in ("id", "key_opt")) for s in SHEETS}
    ids = {s.key: [r[key_pos[s.key]] for r in data[s.key]] for s in SHEETS}  # row position -> id

    wb = Workbook()
    wb.remove(wb.active)
    for spec in SHEETS:
        live = [c for c in spec.cols if c.header]
        rows = []
        for r in data[spec.key]:
            row = []
            for col, v in zip(spec.cols, r):
                if col.kind == "derived":
                    continue
                if col.kind == "ids":
                    v = ";".join(ids[col.target][i] for i in v)
                elif col.kind == "ref":
                    v = ids[col.target][v] if v >= 0 else ""
                elif col.kind == "date" and v:
                    v = date.fromisoformat(v)
                row.append(None if v in ("", None) else v)
            rows.append(row)
        add_table(wb, spec.sheet, spec.table, [c.header for c in live], rows,
                  {i for i, c in enumerate(live) if c.kind in ("id", "key_opt", "ref", "ids")},
                  {i for i, c in enumerate(live) if c.kind == "date"})

    add_table(wb, ACTION_LINKS.sheet, ACTION_LINKS.table, [c.header for c in ACTION_LINKS.cols], data["AL"], {0, 1})
    add_table(wb, META_SHEET, META_TABLE, ["Key", "Value"], [[k, payload[k]] for k in META_KEYS])
    snaps = [[d, pid] for d, pids in payload["burnup_history"].items() for pid in pids]
    add_table(wb, BURNUP_SHEET, BURNUP_TABLE, list(BURNUP_HEADERS), snaps, {1})

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
