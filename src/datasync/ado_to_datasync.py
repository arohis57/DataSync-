"""Turn the raw ADO pull (see ado_to_excel.pull_all) into the DataSync workbook.

Writes the sheets and tables described in docs/data_model.md, using the same layout as
workbook_writer.py. Pure functions over the ``raw`` dict, so it runs from the saved cache
without touching ADO.

Fields the ADO pull does not carry yet (marked ASSUMPTION or TODO below) are left blank or
derived by a stated rule; confirm them against the ADO query logic. Rows the contract has no
place for (RAID Decisions) are dropped and counted in the log.
"""
from __future__ import annotations

import datetime
import re
from pathlib import Path

from openpyxl import Workbook

try:
    from . import ado_to_excel as base
    from .schema import ACTION_LINKS, META_KEYS, META_SHEET, META_TABLE, SHEETS
    from .workbook_writer import add_table
except ImportError:  # run as a plain script from src/datasync
    import ado_to_excel as base
    from schema import ACTION_LINKS, META_KEYS, META_SHEET, META_TABLE, SHEETS
    from workbook_writer import add_table

STREAMS = {"O2C", "S2P", "D2D", "M2O", "R2R"}
RAID_TYPE = {"Action": "A", "Risk": "R", "Issue": "I"}
PBI_STAGES = (("CRD", "CRD"), ("CUT", "CUT"), ("CONFIG & CUT", "CUT"), ("FDD", "FDD"), ("FUT", "FUT"))
COMPANY_TAG = re.compile(r"\[([^\]]+)\]\s*$")


# ----------------------------------------------------------------- derived fields
def workstream(area_path: str) -> str:
    """First segment under GFPhoenix. ASSUMPTION: OCM from Change_OCM, anything else not a stream is Cross-stream."""
    seg = (area_path or "").split("\\")[1:2]
    seg = seg[0].strip() if seg else ""
    if not seg:
        return "Unassigned"
    if seg in STREAMS:
        return seg
    if seg.upper().startswith("CHANGE_OCM") or seg.upper() == "OCM":
        return "OCM"
    return "Cross-stream"


def org_of(person: str | None) -> str:
    """ASSUMPTION: the org is the trailing '[GF]'-style company tag on the person's name; blank when there is none."""
    m = COMPANY_TAG.search(person or "")
    return m.group(1).strip() if m else ""


def pbi_stage(title: str) -> str:
    t = (title or "").strip().upper()
    for prefix, stage in PBI_STAGES:
        if t.startswith(prefix):
            return stage
    return "BUILD"


def strip_number(title: str) -> str:
    return re.sub(r"^\s*\d+(?:\.\s*\d+)*\s*[-.:]?\s*", "", title or "").strip()


def level_label(bph: dict | None) -> str:
    if not bph:
        return ""
    return "%s - %s" % (base.num_of(bph.get("Title")), strip_number(bph.get("Title"))) if base.num_of(bph.get("Title")) else (bph.get("Title") or "")


def ids_of(links, wid, allowed: dict) -> str:
    """';'-joined ids linked to wid that exist in allowed, in numeric order."""
    return ";".join(sorted((n for n in links.neighbours(wid) if n in allowed), key=int))


def d(v):
    """Datetime from the row builders -> date, for the workbook."""
    return v.date() if isinstance(v, datetime.datetime) else v


# ----------------------------------------------------------------- sheet builders
def epic_rows(raw: dict, eps: list[dict]) -> list[list]:
    """One row per L4: a matched Epic, a BPH L4 with no Epic (BPH Only), or an Epic with no BPH L4."""
    bph = {str(b["WorkItemId"]): b for b in raw["bphs"]}
    parent, children = {}, {}
    for b in raw["bphs"]:
        p = str(b["WorkItemId"])
        for l in b.get("Links") or ():
            t = l.get("TargetWorkItem") or {}
            c = str(t.get("WorkItemId"))
            if l.get("LinkTypeName") == "Child" and t.get("WorkItemType") == "BPH" and c in bph:
                parent[c] = p
                children.setdefault(p, []).append(c)

    def chain(i):
        path = [i]
        while path[0] in parent and len(path) < 10:
            path.insert(0, parent[path[0]])
        return path

    by_num = {}
    for e in eps:
        n = base.num_of(e.get("Title"))
        if n:
            by_num.setdefault(base.first4(n), []).append(e)

    rows, used = [], set()

    def row(epic, l4_bph, path):
        l1, l2, l3 = (bph.get(path[k]) if k < len(path) else None for k in range(3))
        title = (epic or {}).get("Title") or (l4_bph or {}).get("Title")
        number = base.num_of(title) or base.num_of((l4_bph or {}).get("Title")) or ""
        return [number, strip_number(title), level_label(l2),
                workstream(base.area_of(epic)) if epic else "Unassigned",
                str(epic["WorkItemId"]) if epic else None, 0 if epic else 1, level_label(l1), level_label(l3)]

    for i, b in bph.items():
        path = chain(i)
        if len(path) != 4:
            continue
        matches = [e for e in by_num.get(base.num_of(b.get("Title")), []) if e["WorkItemId"] not in used]
        for e in matches:
            used.add(e["WorkItemId"])
            rows.append(row(e, b, path))
        if not matches:
            rows.append(row(None, b, path))
    for e in eps:  # Epics with no BPH L4 at all
        if e["WorkItemId"] not in used:
            rows.append(row(e, None, []))
    rows.sort(key=lambda r: (base.num_key(r[0]), r[4] or ""))
    return rows


def build_sheets(raw: dict, asof: str) -> dict[str, list[list]]:
    wss, pbis, links = base.apply_exclusions(raw)
    eps = sorted(raw["eps"], key=lambda x: x["WorkItemId"])
    epic_ids = {str(e["WorkItemId"]): e for e in eps}
    ws_ids = {str(w["WorkItemId"]): w for w in wss}
    cfg_ids = {str(c["WorkItemId"]): c for c in raw["cfgs"]}
    pbi_ids_with = set()  # build items that have at least one PBI
    for p in pbis:
        pbi_ids_with.update(n for n in links.neighbours(p["WorkItemId"]) if n in cfg_ids)
    ac = raw.get("ac", {})
    out: dict[str, list[list]] = {}

    out["E"] = epic_rows(raw, eps)

    out["W"] = [[str(w["WorkItemId"]), workstream(base.area_of(w)), w.get("Title"), w.get("State"),
                 w.get("Custom_ComplexTopic"), d(base.date_only(w.get("TargetDate"))),
                 ids_of(links, w["WorkItemId"], epic_ids)] for w in sorted(wss, key=lambda x: x["WorkItemId"])]

    # TODO: Approved Date, Decision Type, Date Flag / Check Date need ADO fields the pull does not fetch yet.
    out["K"] = [[str(k["WorkItemId"]), workstream(base.area_of(k)), k.get("Title"), k.get("State"),
                 d(base.date_only(k.get("TargetDate"))), ids_of(links, k["WorkItemId"], epic_ids),
                 base.user_of(k), "", None, "", None, "", org_of(base.user_of(k))]
                for k in sorted(raw["kdds"], key=lambda x: x["WorkItemId"])]

    # ASSUMPTION: Origin Epic IDs = Epic IDs; Via Workshop = 1 when the requirement links to no Epic.
    out["Q"] = []
    for r in sorted(raw["reqs"], key=lambda x: x["WorkItemId"]):
        eids = ids_of(links, r["WorkItemId"], epic_ids)
        out["Q"].append([str(r["WorkItemId"]), workstream(base.area_of(r)), r.get("State"),
                         r.get("Custom_SolutionDispositionPHX"), d(base.date_only(r.get("TargetDate"))),
                         eids, ids_of(links, r["WorkItemId"], cfg_ids), r.get("Title"),
                         0 if any(n in ws_ids for n in links.neighbours(r["WorkItemId"])) else 1,
                         base.user_of(r), eids, 0 if eids else 1, org_of(base.user_of(r))])

    out["C"] = []
    for c in sorted(raw["cfgs"], key=lambda x: x["WorkItemId"]):
        cid = str(c["WorkItemId"])
        flags = ("P" if cid in pbi_ids_with else "") + ("E" if any(n in epic_ids for n in links.neighbours(cid)) else "")
        out["C"].append([cid, "C" if c.get("WorkItemType") == "Configuration" else "R", c.get("Title"), c.get("State"),
                         workstream(base.area_of(c)), d(base.date_only(c.get("TargetDate"))),
                         base.tags_of(c), flags, 0])

    out["P"] = []
    for p in sorted(pbis, key=lambda x: x["WorkItemId"]):
        parents = sorted((n for n in links.neighbours(p["WorkItemId"]) if n in cfg_ids), key=int)
        out["P"].append([str(p["WorkItemId"]), pbi_stage(p.get("Title")), p.get("State"), workstream(base.area_of(p)),
                         d(base.date_only(p.get("TargetDate"))), parents[0] if parents else None, p.get("Title"),
                         base.user_of(p), p.get("Custom_ResponsibleParty") or ""])  # ASSUMPTION: Org = Responsible Party

    raids = [x for x in raw["raids"] if x.get("WorkItemType") in RAID_TYPE]
    base.log("RAID Decisions left out (no contract type): %d" % (len(raw["raids"]) - len(raids)))
    out["R"] = []
    for x in sorted(raids, key=lambda x: x["WorkItemId"]):
        tags = base.tags_of(x)
        out["R"].append([str(x["WorkItemId"]), RAID_TYPE[x["WorkItemType"]], x.get("Title"), x.get("State"),
                         workstream(base.area_of(x)), d(base.date_only(x.get("TargetDate"))),
                         base.user_of(x), org_of(base.user_of(x)),  # ASSUMPTION: Firm = company tag on the assignee
                         x.get("Custom_SubTypePHX") or "", tags, x["WorkItemType"]])

    raid_actions = {str(x["WorkItemId"]) for x in raids if x["WorkItemType"] == "Action"}
    out["AL"] = [[a, w] for a in sorted(raid_actions, key=int) for w in sorted(links.neighbours(a), key=int) if w in ws_ids]
    out["META"] = [[k, v] for k, v in zip(META_KEYS, (asof, asof))]  # ASSUMPTION: the export is run on the as-of date
    return out


# ----------------------------------------------------------------- workbook
def write_datasync_workbook(raw: dict, path: str | Path, asof: str | None = None) -> Path:
    asof = asof or datetime.date.today().isoformat()
    datetime.date.fromisoformat(asof)  # fail early on a bad --asof
    sheets = build_sheets(raw, asof)
    wb = Workbook()
    wb.remove(wb.active)
    for spec in SHEETS:
        live = [c for c in spec.cols if c.header]
        rows = [[v for col, v in zip(spec.cols, r) if col.kind != "derived"] for r in sheets[spec.key]]
        rows = [[None if v in ("", None) else v for v in r] for r in rows]
        add_table(wb, spec.sheet, spec.table, [c.header for c in live], rows,
                  {i for i, c in enumerate(live) if c.kind in ("id", "key_opt", "ref", "ids")},
                  {i for i, c in enumerate(live) if c.kind == "date"})
        base.log("  %s: %d rows" % (spec.sheet, len(rows)))
    add_table(wb, ACTION_LINKS.sheet, ACTION_LINKS.table, [c.header for c in ACTION_LINKS.cols], sheets["AL"], {0, 1})
    add_table(wb, META_SHEET, META_TABLE, ["Key", "Value"], sheets["META"])
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    base.log("Saved " + str(path))
    return path
