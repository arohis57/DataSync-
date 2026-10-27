import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from datasync.excel_to_model import DEFAULT_CONFIG, build_payload, main
from datasync.schema import ACTION_LINKS, SHEETS
from datasync.workbook_writer import write_workbook

ROOT = Path(__file__).resolve().parents[1]
REAL_JSON = ROOT / "src" / "cockpit" / "cockpit_data.json"


def tiny_payload() -> dict:
    config = json.loads(DEFAULT_CONFIG.read_text(encoding="utf-8"))
    data = {
        "E": [["1.1.1.1", "L4 one", "L2", "O2C", "100", 0, "L1", "L3"],
              ["1.1.1.2", "L4 two", "L2", "O2C", "101", 0, "L1", "L3"]],
        "W": [["200", "O2C", "Initiate - x", "New", "", "2026-10-12", [0, 1]]],
        "K": [["300", "O2C", "kdd", "New", "2026-10-15", [0], "Doe, J", "", None, "", None, "Strategic Decisions", "GF"]],
        "Q": [["400", "O2C", "New", "GAP", "2026-10-01", [0], [0], "req", 0, "Doe, J", [0], 0, "GF"]],
        "C": [["500", "C", "cfg", "New", "O2C", None, "", "PE", 0]],
        "P": [["600", "BUILD", "New", "O2C", "2026-11-01", 0, "pbi", "Doe, J", "", "GF"],
              ["601", "FUT", "New", "O2C", None, -1, "orphan pbi", "", "N", "GF"]],
        "R": [["700", "A", "act", "New", "O2C", "2026-10-01", "Doe, J", "GF", "OM", "", "Action"]],
        "AL": [["700", "200"], ["999", "200"]],
    }
    return {**config, "asof": "2026-10-25", "burnup_latest_export": "2026-10-26",
            "burnup_history": {"2026-10-18": ["600"]}, "data": data}


@pytest.fixture
def tiny(tmp_path):
    payload = tiny_payload()
    path = tmp_path / "tiny.xlsx"
    write_workbook(payload, path)
    return payload, path


def edit(path: Path, sheet: str, header: str, row: int, value) -> None:
    """Set one cell (row 1 = first data row) of a sheet, found by its header."""
    wb = load_workbook(path)
    ws = wb[sheet]
    col = next(c.column for c in ws[1] if c.value == header)
    ws.cell(row + 1, col).value = value
    wb.save(path)


def messages(rep, kind="warnings"):
    return " | ".join(w["message"] for w in getattr(rep, kind))


def test_tiny_roundtrip(tiny):
    payload, path = tiny
    result, rep = build_payload(path)
    assert not rep.errors and not rep.warnings
    assert result == payload


def test_pbi_without_parent_gets_link_reason(tiny):
    _, path = tiny
    result, _ = build_payload(path)
    orphan = result["data"]["P"][1]
    assert orphan[5] == -1 and orphan[8] == "N"
    assert result["data"]["P"][0][8] == ""


def test_numeric_ids_and_datetime_cells_are_normalised(tiny):
    _, path = tiny
    edit(path, "Workshops", "Epic IDs", 1, 100.0)  # Excel stores a lone number as a float
    edit(path, "Requirements", "Target Date", 1, "2026-10-01T00:00:00")
    result, rep = build_payload(path)
    assert not rep.errors and not rep.warnings
    assert result["data"]["W"][0][6] == [0]
    assert result["data"]["Q"][0][4] == "2026-10-01"


def test_dangling_link_is_dropped_with_warning(tiny):
    _, path = tiny
    edit(path, "Workshops", "Epic IDs", 1, "100;999")
    result, rep = build_payload(path)
    assert result["data"]["W"][0][6] == [0]
    assert "999" in messages(rep)


def test_duplicate_link_is_dropped_with_warning(tiny):
    _, path = tiny
    edit(path, "Workshops", "Epic IDs", 1, "100;101;100")
    result, rep = build_payload(path)
    assert result["data"]["W"][0][6] == [0, 1]
    assert "twice" in messages(rep)


def test_bad_date_is_cleared_with_warning(tiny):
    _, path = tiny
    edit(path, "KDDs", "Target Date", 1, "next Tuesday")
    result, rep = build_payload(path)
    assert result["data"]["K"][0][4] is None
    assert "not a date" in messages(rep)


def test_pbi_with_unknown_parent_becomes_orphan(tiny):
    _, path = tiny
    edit(path, "PBIs", "Build Item ID", 1, "424242")
    result, rep = build_payload(path)
    assert result["data"]["P"][0][5] == -1 and result["data"]["P"][0][8] == "N"
    assert "424242" in messages(rep)


def test_action_link_to_unknown_workshop_warns_but_is_kept(tiny):
    _, path = tiny
    edit(path, ACTION_LINKS.sheet, "Workshop ID", 1, "555")
    result, rep = build_payload(path)
    assert ["700", "555"] in result["data"]["AL"]
    assert "555" in messages(rep)


def test_duplicate_id_is_an_error_and_blocks_output(tiny):
    _, path = tiny
    edit(path, "L4_Epics", "Epic ID", 2, "100")
    result, rep = build_payload(path)
    assert result is None
    assert "duplicate id 100" in messages(rep, "errors")


def test_blank_required_id_is_an_error(tiny):
    _, path = tiny
    edit(path, "Workshops", "Workshop ID", 1, None)  # other cells still filled, so the row is kept
    result, rep = build_payload(path)
    assert result is None and "blank id" in messages(rep, "errors")


def test_missing_column_is_an_error(tiny):
    _, path = tiny
    edit(path, "PBIs", "Stage", 0, "Phase")  # row 0 is the header row: rename the column away
    result, rep = build_payload(path)
    assert result is None
    assert "Stage" in messages(rep, "errors")


def test_missing_table_is_an_error(tiny):
    _, path = tiny
    wb = load_workbook(path)
    wb["RAID"].title = "RAID_old"
    wb.save(path)
    result, rep = build_payload(path)
    assert result is None
    assert "RAID" in messages(rep, "errors")


def test_burnup_sheet_is_optional(tiny):
    payload, path = tiny
    wb = load_workbook(path)
    del wb["Burnup_History"]
    wb.save(path)
    result, rep = build_payload(path)
    assert not rep.errors and result["burnup_history"] == {}


def test_missing_meta_setting_is_an_error(tiny):
    _, path = tiny
    edit(path, "Meta", "Key", 1, "something_else")
    result, rep = build_payload(path)
    assert result is None
    assert "asof" in messages(rep, "errors")


def test_cli_writes_payload_and_report(tiny, tmp_path):
    payload, path = tiny
    out = tmp_path / "out" / "cockpit_data.json"
    assert main([str(path), "-o", str(out)]) == 0
    assert json.loads(out.read_text(encoding="utf-8")) == payload
    assert json.loads((out.parent / "validation_report.json").read_text())["ok"] is True


def test_cli_writes_nothing_on_error(tiny, tmp_path):
    _, path = tiny
    edit(path, "L4_Epics", "Epic ID", 2, "100")
    out = tmp_path / "out" / "cockpit_data.json"
    assert main([str(path), "-o", str(out)]) == 1
    assert not out.exists()
    assert json.loads((out.parent / "validation_report.json").read_text())["ok"] is False


@pytest.mark.skipif(not REAL_JSON.exists(), reason="reference cockpit_data.json not present")
def test_real_data_roundtrips_exactly(tmp_path):
    original = json.loads(REAL_JSON.read_text(encoding="utf-8"))
    path = tmp_path / "real.xlsx"
    write_workbook(original, path)
    result, rep = build_payload(path)
    assert not rep.errors and not rep.warnings
    assert result == original


def test_data_model_doc_lists_every_column():
    doc = (ROOT / "docs" / "data_model.md").read_text(encoding="utf-8")
    for spec in (*SHEETS, ACTION_LINKS):
        assert f"`{spec.sheet}`" in doc and f"`{spec.table}`" in doc, spec.sheet
        for col in spec.cols:
            if col.header:
                assert f"`{col.header}`" in doc, f"{spec.sheet}.{col.header} missing from docs/data_model.md"
